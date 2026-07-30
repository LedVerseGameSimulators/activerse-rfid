"""Excel onboard for companies / groups / players."""
from io import BytesIO
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook

from . import players
from .database import Database

HEADERS = [
    "company_name",
    "group_name",
    "player_name",
    "phone",
    "is_team_leader",
    "email",
    "age",
]


def _truthy_leader(val) -> bool:
    if val is None:
        return False
    s = str(val).strip().lower()
    return s in ("y", "yes", "true", "1", "leader")


def build_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "import"
    ws.append(HEADERS)
    ws.append(["Acme Corp", "Team Alpha", "Ada Lovelace", "9990001111", "Y", "ada@acme.test", 30])
    ws.append(["Acme Corp", "Team Alpha", "Alan Turing", "9990002222", "N", "", ""])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_excel(db: Database, upload: UploadFile) -> dict:
    raw = upload.file.read()
    try:
        wb = load_workbook(BytesIO(raw), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"Invalid Excel file: {e}")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "Empty spreadsheet")
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    for h in ("company_name", "group_name", "player_name", "phone", "is_team_leader"):
        if h not in header:
            raise HTTPException(400, f"Missing column: {h}")

    idx = {name: header.index(name) for name in header}

    report: dict[str, Any] = {
        "created_companies": 0,
        "created_groups": 0,
        "created_players": 0,
        "linked_players": 0,
        "leaders_set": 0,
        "errors": [],
    }
    # group_key -> list of (row_num, player_id, is_leader)
    pending_leaders: dict[tuple, list] = {}

    for i, row in enumerate(rows[1:], start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        def cell(name, default=""):
            j = idx.get(name)
            if j is None or j >= len(row) or row[j] is None:
                return default
            return row[j]

        company_name = str(cell("company_name")).strip()
        group_name = str(cell("group_name")).strip()
        player_name = str(cell("player_name")).strip()
        phone = str(cell("phone")).strip()
        if phone.endswith(".0"):  # excel numeric phones
            phone = phone[:-2]
        is_leader = _truthy_leader(cell("is_team_leader"))
        email = str(cell("email")).strip() or None
        age_raw = cell("age", None)
        age = None
        if age_raw not in (None, ""):
            try:
                age = int(float(age_raw))
            except (TypeError, ValueError):
                report["errors"].append({"row": i, "message": f"Invalid age: {age_raw}"})
                continue

        if not company_name or not group_name or not player_name or not phone:
            report["errors"].append({"row": i, "message": "company_name, group_name, player_name, phone required"})
            continue

        try:
            company = db.get_company_by_name(company_name)
            if not company:
                company = db.create_company(company_name)
                report["created_companies"] += 1

            group = db.get_group_by_company_name(company["id"], group_name)
            if not group:
                group = db.create_group(company["id"], group_name)
                report["created_groups"] += 1

            existing = db.search_custom_by_field("phone_num", phone)
            if existing:
                pid = existing[0]["custom_id"]
                report["linked_players"] += 1
            else:
                created = players.register_player(db, player_name, phone, email, age)
                pid = created["id"]
                report["created_players"] += 1

            db.add_group_member_atomic(group["id"], pid)
            key = (company["id"], group["id"])
            pending_leaders.setdefault(key, []).append((i, pid, is_leader))
        except Exception as e:
            report["errors"].append({"row": i, "message": str(e)})

    for (company_id, group_id), entries in pending_leaders.items():
        leaders = [e for e in entries if e[2]]
        if len(leaders) == 0:
            report["errors"].append({
                "row": entries[0][0],
                "message": f"Group {group_id}: no team leader marked",
            })
            continue
        if len(leaders) > 1:
            report["errors"].append({
                "row": leaders[0][0],
                "message": f"Group {group_id}: multiple team leaders marked",
            })
            continue
        db.set_group_leader(group_id, leaders[0][1])
        report["leaders_set"] += 1

    return report


HEADERS_COMPANY_SCOPED = ["group_name", "player_name", "phone", "is_team_leader", "email", "age"]


def build_template_bytes_company_scoped() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "import"
    ws.append(HEADERS_COMPANY_SCOPED)
    ws.append(["Team Alpha", "Ada Lovelace", "9990001111", "Y", "ada@acme.test", 30])
    ws.append(["Team Alpha", "Alan Turing", "9990002222", "N", "", ""])
    ws.append(["", "Solo Player (no team yet)", "9990003333", "N", "", ""])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_excel_for_company(db: Database, company_id: int, upload: UploadFile) -> dict:
    """Company is fixed by company_id (URL) — never created/looked-up per row.
    Any company_name column in the sheet is ignored entirely. group_name is
    optional: blank means company-only (no group). Idempotent on re-upload —
    every lookup is get-or-create keyed on fixed identifiers."""
    if not db.get_company(company_id):
        raise HTTPException(404, "Company not found")

    raw = upload.file.read()
    try:
        wb = load_workbook(BytesIO(raw), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"Invalid Excel file: {e}")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "Empty spreadsheet")
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    for h in ("player_name", "phone", "is_team_leader"):
        if h not in header:
            raise HTTPException(400, f"Missing column: {h}")

    idx = {name: header.index(name) for name in header}

    report: dict[str, Any] = {
        "created_groups": 0,
        "created_players": 0,
        "linked_players": 0,
        "company_members_added": 0,
        "leaders_set": 0,
        "errors": [],
    }
    pending_leaders: dict[int, list] = {}

    for i, row in enumerate(rows[1:], start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        def cell(name, default=""):
            j = idx.get(name)
            if j is None or j >= len(row) or row[j] is None:
                return default
            return row[j]

        group_name = str(cell("group_name")).strip()
        player_name = str(cell("player_name")).strip()
        phone = str(cell("phone")).strip()
        if phone.endswith(".0"):  # excel numeric phones
            phone = phone[:-2]
        is_leader = _truthy_leader(cell("is_team_leader"))
        email = str(cell("email")).strip() or None
        age_raw = cell("age", None)
        age = None
        if age_raw not in (None, ""):
            try:
                age = int(float(age_raw))
            except (TypeError, ValueError):
                report["errors"].append({"row": i, "message": f"Invalid age: {age_raw}"})
                continue

        if not player_name or not phone:
            report["errors"].append({"row": i, "message": "player_name, phone required"})
            continue

        try:
            existing = db.search_custom_by_field("phone_num", phone)
            if existing:
                pid = existing[0]["custom_id"]
                report["linked_players"] += 1
            else:
                created = players.register_player(db, player_name, phone, email, age)
                pid = created["id"]
                report["created_players"] += 1

            db.add_company_member_atomic(company_id, pid)
            report["company_members_added"] += 1

            if group_name:
                group = db.get_group_by_company_name(company_id, group_name)
                if not group:
                    group = db.create_group(company_id, group_name)
                    report["created_groups"] += 1
                db.add_group_member_atomic(group["id"], pid)
                pending_leaders.setdefault(group["id"], []).append((i, pid, is_leader))
        except ValueError as e:
            report["errors"].append({"row": i, "message": str(e)})
        except Exception as e:
            report["errors"].append({"row": i, "message": str(e)})

    for group_id, entries in pending_leaders.items():
        leaders = [e for e in entries if e[2]]
        if len(leaders) == 1:
            db.set_group_leader(group_id, leaders[0][1])
            report["leaders_set"] += 1
        elif len(leaders) > 1:
            report["errors"].append({
                "row": leaders[0][0],
                "message": f"Group {group_id}: multiple team leaders marked",
            })

    return report
